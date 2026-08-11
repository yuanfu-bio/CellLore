"""Read a completed run's artifacts into an auditable curation catalog.

Shaped after ``04_thinking/agent.xlsx``, the curator's own reference for what
"done" looks like: one row per study, one per pair unit, one per file, one
per open audit/hold item. Every value here is read off an existing artifact
or claim -- nothing is invented for presentation. Where the pipeline has not
yet produced a value (a genome build nobody declared, a priority no
acquisition plan ran to assign), the cell is left blank rather than guessed.

Internal state uses a much finer-grained lifecycle (``BundleStatus`` has
twenty-odd values; ``PreflightBundleReport.decision`` only ever lands on
three of them pre-acquisition). This module collapses that onto exactly the
three top-level states the catalog is allowed to show --
``READY``/``REVIEW``/``HOLD`` -- and pushes every nuance into
``catalog_status_detail``/``next_action``/the audit row, never a fourth
top-level status.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from typing import Any

from celllore.application.agents.shared import latest_artifact
from celllore.application.collaboration import AgentArtifact, ArtifactKind
from celllore.application.dataset_layout import readable_pair_unit_id
from celllore.domain.curation.profiles import profile_by_id
from celllore.domain.lifecycle import (
    AcquisitionBundle,
    AuditQueueAction,
    AuditTask,
    BundleStatus,
    EvidenceClaim,
    PreflightBundleReport,
    SemanticBundle,
)
from celllore.domain.models import (
    DataAsset,
    DataRole,
    DatasetCandidate,
    DiscoveryRunResult,
    SourceRecord,
)
from celllore.domain.pair_units import ProposedPairUnit

CATALOG_STATUS_BY_DECISION: dict[BundleStatus, str] = {
    BundleStatus.PREFLIGHT_READY: "READY",
    BundleStatus.NEEDS_EVIDENCE: "REVIEW",
    BundleStatus.REJECTED: "HOLD",
}

_NEXT_ACTION_BY_QUEUE_ACTION: dict[AuditQueueAction, str] = {
    AuditQueueAction.VERIFY_THEN_ARCHIVE: "Verify the flagged evidence, then archive.",
    AuditQueueAction.DEFERRED: "Defer until the missing file or role can be located.",
    AuditQueueAction.EXCLUDED: "Excluded from this profile's corpus.",
    AuditQueueAction.REPROCESS: "Needs reprocessing (re-download or re-run a pipeline step).",
    AuditQueueAction.PARTIAL_ARCHIVE: "Archive the resolved pair units; hold the rest.",
}

_FAILURE_HANDLING_BY_QUEUE_ACTION: dict[AuditQueueAction, str] = {
    AuditQueueAction.VERIFY_THEN_ARCHIVE: (
        "If verification fails, keep this pair unit at REVIEW rather than "
        "promoting it to READY."
    ),
    AuditQueueAction.DEFERRED: (
        "Stays deferred; re-check once source metadata is re-fetched or the "
        "missing file appears."
    ),
    AuditQueueAction.EXCLUDED: (
        "Stays excluded from this corpus; a future run can re-admit it if "
        "new evidence resolves the gap."
    ),
    AuditQueueAction.REPROCESS: (
        "If reprocessing does not resolve it, escalate to manual review "
        "rather than retrying indefinitely."
    ),
    AuditQueueAction.PARTIAL_ARCHIVE: (
        "The unresolved pair units stay at REVIEW/HOLD independently of the "
        "ones archived; one bad pair unit must not block its siblings."
    ),
}

_ATAC_ROLES = {DataRole.FRAGMENTS, DataRole.FRAGMENT_INDEX, DataRole.PEAKS, DataRole.PEAK_MATRIX}
_RNA_ROLES = {DataRole.RNA_MATRIX, DataRole.STANDARD_OBJECT}
_METADATA_ROLES = {
    DataRole.CELL_METADATA, DataRole.BARCODE_LINKAGE, DataRole.SAMPLE_METADATA,
    DataRole.DONOR_METADATA, DataRole.QC_METRICS,
}

_GSM_RE = re.compile(r"GSM\d+")

_DOWNSTREAM_ACTION_BY_ROLE: dict[DataRole, str] = {
    DataRole.FRAGMENTS: "Verify barcode-level fragment format and build a tabix index.",
    DataRole.FRAGMENT_INDEX: "Confirm the index matches its fragment file before use.",
    DataRole.RNA_MATRIX: "Verify the matrix holds unnormalized integer counts, not normalized values.",
    DataRole.STANDARD_OBJECT: "Open the object and confirm it carries a raw, unnormalized count layer.",
    DataRole.CELL_METADATA: "Cross-check barcodes against the paired modality's own barcode set.",
    DataRole.BARCODE_LINKAGE: "Apply this translation before intersecting RNA/ATAC barcodes.",
    DataRole.ARCHIVE: "Open the archive and enumerate its members; contents are unknown, not absent.",
    DataRole.PEAKS: "Not required for the current profile; keep for downstream peak-calling reuse.",
    DataRole.QC_METRICS: "Reference only; not part of the current profile's required roles.",
}


@dataclass
class CurationCatalog:
    studies: list[dict[str, Any]] = field(default_factory=list)
    pair_units: list[dict[str, Any]] = field(default_factory=list)
    files: list[dict[str, Any]] = field(default_factory=list)
    audit_hold: list[dict[str, Any]] = field(default_factory=list)
    policy: list[dict[str, Any]] = field(default_factory=list)

    def sheets(self) -> list[tuple[str, list[dict[str, Any]]]]:
        return [
            ("01_Studies", self.studies),
            ("02_Pair_units", self.pair_units),
            ("03_Files", self.files),
            ("04_Audit_Hold", self.audit_hold),
            ("05_Policy", self.policy),
        ]


def write_curation_catalog_xlsx(catalog: CurationCatalog, path: str) -> None:
    """Write the catalog to an .xlsx with one sheet per level, matching
    ``04_thinking/agent.xlsx``'s own sheet names."""
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in catalog.sheets():
        sheet = workbook.create_sheet(name)
        if not rows:
            continue
        headers = list(rows[0].keys())
        sheet.append(headers)
        for row in rows:
            sheet.append(["" if row[key] is None else str(row[key]) for key in headers])
            for cell in sheet[sheet.max_row]:
                value = cell.value
                if isinstance(value, str) and value.startswith(
                    ("https://", "http://", "ftp://")
                ):
                    cell.hyperlink = value
                    cell.style = "Hyperlink"
    workbook.save(path)


def build_curation_catalog(
    run_id: str, artifacts: list[AgentArtifact]
) -> CurationCatalog:
    """Build the catalog from one run's artifact list (``ledger.list_artifacts(run_id)``)."""
    discovery_artifact = latest_artifact(artifacts, ArtifactKind.DISCOVERY_RESULT.value)
    known_fields = set(DiscoveryRunResult.model_fields)
    discovery = (
        DiscoveryRunResult.model_validate(
            _without_legacy_modalities(
                {k: v for k, v in discovery_artifact.payload.items() if k in known_fields}
            )
        )
        if discovery_artifact is not None
        else None
    )
    candidates_by_id = {
        item.candidate_id: item for item in (discovery.candidates if discovery else [])
    }
    source_records_by_id = {
        item.record_id: item for item in (discovery.source_records if discovery else [])
    }

    inventory_artifact = latest_artifact(artifacts, ArtifactKind.ASSET_INVENTORY.value)
    payload = inventory_artifact.payload if inventory_artifact is not None else {}
    assets_by_id = {
        item["asset_id"]: DataAsset.model_validate(item)
        for item in payload.get("assets", [])
    }
    bundles = [SemanticBundle.model_validate(item) for item in payload.get("bundles", [])]
    required_layers = {
        layer.value
        for layer in (
            discovery.requirement.required_omics_layers if discovery else []
        )
    }
    if required_layers:
        bundles = [
            bundle
            for bundle in bundles
            if required_layers <= {
                _omics_layer_for_role(asset.role)
                for asset in _bundle_assets(bundle, assets_by_id)
            }
        ]
    proposed_units_by_id = {
        item["proposed_unit_id"]: ProposedPairUnit.model_validate(item)
        for item in payload.get("proposed_pair_units", [])
    }
    pairing_claims_by_unit = {
        item["subject_id"]: EvidenceClaim.model_validate(item)
        for item in payload.get("pairing_claims", [])
    }

    preflight_artifact = latest_artifact(artifacts, ArtifactKind.PREFLIGHT_REPORT.value)
    preflight_payload = preflight_artifact.payload if preflight_artifact is not None else {}
    reports_by_bundle = {
        item["bundle_id"]: PreflightBundleReport.model_validate(item)
        for item in preflight_payload.get("reports", [])
    }
    audit_tasks = [
        AuditTask.model_validate(item) for item in preflight_payload.get("audit_tasks", [])
    ]
    assembly_by_bundle: dict[str, list[str]] = {
        item["subject_id"]: item["value"]
        for item in preflight_payload.get("assembly_claims", [])
        if item.get("predicate") == "resolved_assembly"
    }

    acquisition_artifact = latest_artifact(artifacts, ArtifactKind.ACQUISITION_PLAN.value)
    acquisition_payload = acquisition_artifact.payload if acquisition_artifact is not None else {}
    acquisition_by_bundle = {
        item["bundle_id"]: AcquisitionBundle.model_validate(item)
        for item in acquisition_payload.get("bundles", [])
    }

    bundles_by_candidate: dict[str, list[SemanticBundle]] = {}
    for bundle in bundles:
        bundles_by_candidate.setdefault(bundle.candidate_id, []).append(bundle)

    catalog = CurationCatalog()
    if discovery and discovery.requirement.profile_id:
        try:
            profile = profile_by_id(discovery.requirement.profile_id)
        except KeyError:
            profile = None
        if profile is not None:
            catalog.policy = [
                {"requirement_id": item.requirement_id, "statement": item.statement,
                 "profile_id": profile.profile_id}
                for item in profile.hard_requirements
            ]
    # Discovery is the authoritative study list. A candidate without a
    # reconciled bundle is still a useful curation outcome: silently dropping
    # it makes an incomplete GEO manifest look undiscovered rather than a
    # record that needs a concrete follow-up.
    for candidate_id, candidate in candidates_by_id.items():
        candidate_bundles = bundles_by_candidate.get(candidate_id, [])
        catalog.studies.append(
            _study_row(
                candidate, candidate_bundles, reports_by_bundle, assets_by_id,
                pairing_claims_by_unit, proposed_units_by_id, source_records_by_id,
            )
        )
        if not candidate_bundles:
            catalog.audit_hold.append(_unresolved_study_audit_row(candidate))
        for bundle in candidate_bundles:
            display_pair_unit_id = readable_pair_unit_id(
                bundle,
                proposed_units_by_id.get(bundle.pair_unit_id) if bundle.pair_unit_id else None,
                _bundle_assets(bundle, assets_by_id),
            )
            catalog.pair_units.append(
                _pair_unit_row(
                    candidate, bundle, proposed_units_by_id, pairing_claims_by_unit,
                    reports_by_bundle, assembly_by_bundle, acquisition_by_bundle, assets_by_id,
                    display_pair_unit_id,
                )
            )
            catalog.files.extend(
                _file_rows(
                    candidate, bundle, reports_by_bundle, acquisition_by_bundle, assets_by_id,
                    display_pair_unit_id,
                )
            )

    bundles_by_id = {item.bundle_id: item for item in bundles}
    for task in audit_tasks:
        if task.bundle_id is not None and task.bundle_id not in bundles_by_id:
            continue
        catalog.audit_hold.append(
            _audit_row(task, bundles_by_id, candidates_by_id, acquisition_by_bundle)
        )

    return catalog


def _without_legacy_modalities(value: Any) -> Any:
    """Read pre-OmicsLayer artifacts without weakening current schemas."""
    if isinstance(value, dict):
        return {key: _without_legacy_modalities(item) for key, item in value.items()
                if key not in {"modalities", "acceptable_modalities", "preferred_modalities"}}
    if isinstance(value, list):
        return [_without_legacy_modalities(item) for item in value]
    return value


def _geo_series(candidate: DatasetCandidate) -> str | None:
    for item in candidate.identifiers:
        if item.namespace.casefold() == "geo_series":
            return item.value
    return None


def _study_accession(candidate: DatasetCandidate) -> str:
    return _geo_series(candidate) or candidate.candidate_id


def _repository(candidate: DatasetCandidate) -> str:
    return "GEO" if _geo_series(candidate) is not None else "unknown"


def _public_year(
    candidate: DatasetCandidate,
    source_records_by_id: dict[str, SourceRecord],
) -> int | None:
    """Read the repository/publication year from linked discovery metadata."""
    records = [
        source_records_by_id[record_id]
        for record_id in candidate.source_record_ids
        if record_id in source_records_by_id
    ]
    records.sort(key=lambda item: 0 if item.source_name == "ncbi_geo" else 1)
    preferred_keys = (
        "pdat",
        "pubyear",
        "firstPublicationDate",
        "electronicPublicationDate",
        "dateOfPublication",
    )
    for record in records:
        metadata = record.raw_metadata or {}
        values = [metadata.get(key) for key in preferred_keys]
        journal_info = metadata.get("journalInfo")
        if isinstance(journal_info, dict):
            values.extend(
                journal_info.get(key)
                for key in ("printPublicationDate", "electronicPublicationDate")
            )
        for value in values:
            if match := re.search(r"\b(?:19|20)\d{2}\b", str(value or "")):
                return int(match.group())
    return None


def _geo_url(accession: str) -> str:
    return f"https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc={accession}"


def _method(candidate: DatasetCandidate) -> str:
    terms = [
        term
        for profile in candidate.assay_profiles
        for term in profile.technology_terms
    ]
    if terms:
        return ", ".join(dict.fromkeys(terms))
    labels = [profile.label for profile in candidate.assay_profiles]
    return ", ".join(dict.fromkeys(labels)) or "unknown"


def _human_material(candidate: DatasetCandidate) -> str:
    parts = [*candidate.tissues, *candidate.diseases]
    return ", ".join(dict.fromkeys(parts)) or ", ".join(candidate.organisms) or "unknown"


def _catalog_status(bundle_statuses: list[str]) -> str:
    if not bundle_statuses:
        # No PairUnit means GSM/file-level resolution did not happen. This is
        # a blocking gap for this catalog, not a clean study with zero rows.
        return "HOLD"
    if "READY" in bundle_statuses:
        return "READY"
    if "REVIEW" in bundle_statuses:
        return "REVIEW"
    if "HOLD" in bundle_statuses:
        return "HOLD"
    # No preflight report reached this candidate yet -- not yet assessed is
    # closer to "needs a human look" than to a clean pass.
    return "REVIEW"


def _bundle_status(bundle: SemanticBundle, reports_by_bundle: dict[str, PreflightBundleReport]) -> str:
    report = reports_by_bundle.get(bundle.bundle_id)
    if report is None:
        return "REVIEW"
    return CATALOG_STATUS_BY_DECISION.get(report.decision, "REVIEW")


def _bundle_assets(bundle: SemanticBundle, assets_by_id: dict[str, DataAsset]) -> list[DataAsset]:
    return [assets_by_id[item] for item in bundle.asset_ids if item in assets_by_id]


def _has_role(assets: list[DataAsset], roles: set[DataRole]) -> bool:
    return any(item.role in roles for item in assets)


def _study_row(
    candidate: DatasetCandidate,
    bundles: list[SemanticBundle],
    reports_by_bundle: dict[str, PreflightBundleReport],
    assets_by_id: dict[str, DataAsset],
    pairing_claims_by_unit: dict[str, EvidenceClaim],
    proposed_units_by_id: dict[str, ProposedPairUnit],
    source_records_by_id: dict[str, SourceRecord],
) -> dict[str, Any]:
    accession = _study_accession(candidate)
    statuses = [_bundle_status(item, reports_by_bundle) for item in bundles]
    status_counts = {label: statuses.count(label) for label in ("READY", "REVIEW", "HOLD")}
    all_assets = [asset for bundle in bundles for asset in _bundle_assets(bundle, assets_by_id)]
    pairing_summary = ", ".join(dict.fromkeys(
        f"{readable_pair_unit_id(bundle, proposed_units_by_id.get(bundle.pair_unit_id), _bundle_assets(bundle, assets_by_id))}: {report.decision.value}"
        for bundle in bundles
        if (report := reports_by_bundle.get(bundle.bundle_id)) is not None
        and report.decision != BundleStatus.PREFLIGHT_READY
    ))
    return {
        "study_accession": accession,
        "study_title": candidate.title,
        "repository": _repository(candidate),
        "public_year": _public_year(candidate, source_records_by_id),
        "method": _method(candidate),
        "human_material": _human_material(candidate),
        "assembly": ", ".join(dict.fromkeys(
            build
            for bundle in bundles
            for build in _bundle_assembly(bundle, assets_by_id)
        )) or "unknown",
        "source_url": _geo_url(accession) if _geo_series(candidate) else None,
        "evidence_url": (
            _geo_url(gsm) if (gsm := _first_asset_gsm(all_assets)) else None
        ),
        "catalog_status": _catalog_status(statuses),
        "catalog_status_detail": (
            "No pair unit or exact file manifest was resolved; see 04_Audit_Hold."
            if not bundles
            else f"{status_counts['READY']} READY, {status_counts['REVIEW']} REVIEW, "
            f"{status_counts['HOLD']} HOLD of {len(bundles)} pair unit(s)."
            + (f" Blocked: {pairing_summary}." if pairing_summary else "")
        ),
        "priority": None,
        "next_action": (
            "Fetch/parse GSM metadata and resolve an exact RNA/ATAC file manifest."
            if not bundles
            else "Proceed to download the READY pair unit(s)."
            if status_counts["READY"]
            else "Resolve the REVIEW/HOLD pair units listed in 04_Audit_Hold."
        ),
        "same_cell_evidence": ", ".join(dict.fromkeys(
            f"{readable_pair_unit_id(bundle, proposed_units_by_id.get(bundle.pair_unit_id), _bundle_assets(bundle, assets_by_id))}: {claim.value} ({claim.evidence_class.value})"
            for bundle in bundles
            if bundle.pair_unit_id
            and (claim := pairing_claim_for(bundle, pairing_claims_by_unit))
        )) or None,
        "fragment_evidence": (
            f"{sum(1 for item in all_assets if item.role in _ATAC_ROLES)} ATAC file(s) observed"
            if _has_role(all_assets, _ATAC_ROLES) else "none observed"
        ),
        "rna_count_evidence": (
            f"{sum(1 for item in all_assets if item.role in _RNA_ROLES)} RNA count file(s) observed"
            if _has_role(all_assets, _RNA_ROLES) else "none observed"
        ),
        "metadata_evidence": (
            f"{sum(1 for item in all_assets if item.role in _METADATA_ROLES)} metadata/QC file(s) observed"
            if _has_role(all_assets, _METADATA_ROLES) else "none observed"
        ),
        "notes": "; ".join(candidate.warnings) or None,
        "search_date": None,
        "citation": None,
    }


def pairing_claim_for(
    bundle: SemanticBundle, pairing_claims_by_unit: dict[str, EvidenceClaim]
) -> EvidenceClaim | None:
    if bundle.pair_unit_id is None:
        return None
    return pairing_claims_by_unit.get(bundle.pair_unit_id)


def _bundle_assembly(bundle: SemanticBundle, assets_by_id: dict[str, DataAsset]) -> list[str]:
    return sorted({
        item.genome_build
        for item in _bundle_assets(bundle, assets_by_id)
        if item.genome_build
    })


def _gsm_by_role(assets: list[DataAsset], roles: set[DataRole]) -> list[str]:
    accessions = {
        match.group()
        for item in assets
        if item.role in roles
        for match in [_GSM_RE.search(item.file_name)]
        if match
    }
    return sorted(accessions)


def _first_asset_gsm(assets: list[DataAsset]) -> str | None:
    for item in assets:
        if match := _GSM_RE.search(item.file_name):
            return match.group()
    # GEO URLs contain both a shortened bucket such as ``GSM9729nnn`` and the
    # full accession later in the path. Prefer the longest match if a filename
    # itself did not carry the accession.
    matches = [
        match
        for item in assets
        for location in item.locations
        for match in _GSM_RE.findall(location.url)
    ]
    return max(matches, key=len) if matches else None


def _pair_unit_row(
    candidate: DatasetCandidate,
    bundle: SemanticBundle,
    proposed_units_by_id: dict[str, ProposedPairUnit],
    pairing_claims_by_unit: dict[str, EvidenceClaim],
    reports_by_bundle: dict[str, PreflightBundleReport],
    assembly_by_bundle: dict[str, list[str]],
    acquisition_by_bundle: dict[str, AcquisitionBundle],
    assets_by_id: dict[str, DataAsset],
    display_pair_unit_id: str,
) -> dict[str, Any]:
    accession = _study_accession(candidate)
    unit = proposed_units_by_id.get(bundle.pair_unit_id) if bundle.pair_unit_id else None
    claim = pairing_claim_for(bundle, pairing_claims_by_unit)
    report = reports_by_bundle.get(bundle.bundle_id)
    assets = _bundle_assets(bundle, assets_by_id)
    acquisition = acquisition_by_bundle.get(bundle.bundle_id)
    assembly_values = assembly_by_bundle.get(bundle.bundle_id) or _bundle_assembly(bundle, assets_by_id)
    return {
        "study_accession": accession,
        "pair_unit_id": display_pair_unit_id,
        "sample_condition": (
            (unit.sample_interpretation or unit.grouping_rationale) if unit else None
        ),
        "atac_gsm": ", ".join(_gsm_by_role(assets, _ATAC_ROLES)) or None,
        "rna_gsm": ", ".join(_gsm_by_role(assets, _RNA_ROLES)) or None,
        "assembly": ", ".join(assembly_values) or "unknown",
        "pairing_key": ((unit and unit.member_samples) and " + ".join(unit.member_samples)) or None,
        "pairing_evidence": (
            f"{claim.value} ({claim.evidence_class.value}): {claim.rationale}" if claim else "unverified"
        ),
        "rna_count_scope": (
            "processed count matrix/object" if _has_role(assets, _RNA_ROLES) else "not observed"
        ),
        "metadata_status": (
            "present" if _has_role(assets, _METADATA_ROLES) else "not observed"
        ),
        "manifest_status": report.decision.value if report else "not_yet_assessed",
        "priority": acquisition.priority if acquisition else None,
        "notes": (
            "; ".join(f"{k}: {v}" for k, v in report.rationale.items() if k in {"bundle_schema", "pairing"})
            if report else None
        ),
        "source_url": _geo_url(accession) if _geo_series(candidate) else None,
    }


def _omics_layer_for_role(role: DataRole) -> str:
    if role in _ATAC_ROLES:
        return "chromatin_accessibility"
    if role in _RNA_ROLES:
        return "transcriptome"
    if role is DataRole.PROTEIN_MATRIX:
        return "protein"
    if role in _METADATA_ROLES:
        return "metadata"
    return "other"


def _file_rows(
    candidate: DatasetCandidate,
    bundle: SemanticBundle,
    reports_by_bundle: dict[str, PreflightBundleReport],
    acquisition_by_bundle: dict[str, AcquisitionBundle],
    assets_by_id: dict[str, DataAsset],
    display_pair_unit_id: str,
) -> list[dict[str, Any]]:
    accession = _study_accession(candidate)
    status = _bundle_status(bundle, reports_by_bundle)
    acquisition = acquisition_by_bundle.get(bundle.bundle_id)
    rows = []
    for asset in _bundle_assets(bundle, assets_by_id):
        location = asset.locations[0] if asset.locations else None
        gsm_match = _GSM_RE.search(asset.file_name)
        rows.append({
            "dataset_id": accession,
            "pair_unit_id": display_pair_unit_id,
            "file_accession": gsm_match.group() if gsm_match else None,
            "assembly": asset.genome_build or "unknown",
            "omics_layer": _omics_layer_for_role(asset.role),
            "file_role": asset.role.value,
            "filename": asset.file_name,
            "size_bytes": location.size_bytes if location else None,
            "download_url": location.url if location else None,
            "url_status": (
                f"{location.availability.value}/{location.access_level.value}"
                if location else "unknown"
            ),
            "audit_status": status,
            "priority": acquisition.priority if acquisition else None,
            "expected_downstream_action": _DOWNSTREAM_ACTION_BY_ROLE.get(
                asset.role, "No specific action registered for this role yet."
            ),
        })
    return rows


def _audit_row(
    task: AuditTask,
    bundles_by_id: dict[str, SemanticBundle],
    candidates_by_id: dict[str, DatasetCandidate],
    acquisition_by_bundle: dict[str, AcquisitionBundle],
) -> dict[str, Any]:
    bundle = bundles_by_id.get(task.bundle_id) if task.bundle_id else None
    candidate = candidates_by_id.get(bundle.candidate_id) if bundle else None
    acquisition = acquisition_by_bundle.get(task.bundle_id) if task.bundle_id else None
    accession = _study_accession(candidate) if candidate else None
    return {
        "priority": acquisition.priority if acquisition else None,
        "record_level": "pair_unit" if (bundle and bundle.pair_unit_id) else "study",
        "study_accession": accession,
        "pair_unit_or_sample_scope": task.subject_id,
        "queue_action": task.queue_action.value,
        "blocking_category": task.blocking_category,
        "category_family": task.category_family.value,
        "blocking_reason_detail": task.rationale,
        "audit_method": task.audit_method or None,
        "audit_steps": task.audit_steps or None,
        "pass_criteria": task.pass_criteria or None,
        "failure_handling": _FAILURE_HANDLING_BY_QUEUE_ACTION.get(task.queue_action),
        "next_action": _NEXT_ACTION_BY_QUEUE_ACTION.get(task.queue_action, "Review manually."),
        "source_url": _geo_url(accession) if accession else None,
        "audit_date": task.created_at.date().isoformat(),
        "raw_status_code": task.blocking_category,
    }

def _unresolved_study_audit_row(candidate: DatasetCandidate) -> dict[str, Any]:
    """Represent a discovered study with no curatable unit as an actionable hold."""
    accession = _study_accession(candidate)
    return {
        "priority": None,
        "record_level": "study",
        "study_accession": accession,
        "pair_unit_or_sample_scope": None,
        "queue_action": AuditQueueAction.DEFERRED.value,
        "blocking_category": "pair_unit:unresolved",
        "category_family": "inventory",
        "blocking_reason_detail": (
            "The study was discovered, but no GSM-resolved pair unit and exact "
            "RNA/ATAC file manifest were produced."
        ),
        "audit_method": "GEO GSM metadata and supplementary-file manifest review",
        "audit_steps": (
            "Fetch the GSE's GSM records; map each supplementary file to its GSM; "
            "resolve RNA count and ATAC fragment candidates into PairUnits."
        ),
        "pass_criteria": (
            "At least one PairUnit has traceable GSM membership plus exact RNA and "
            "ATAC file URLs."
        ),
        "failure_handling": _FAILURE_HANDLING_BY_QUEUE_ACTION[AuditQueueAction.DEFERRED],
        "next_action": "Fetch/parse GSM metadata and resolve an exact RNA/ATAC file manifest.",
        "source_url": _geo_url(accession) if _geo_series(candidate) else None,
        "audit_date": None,
        "raw_status_code": "pair_unit:unresolved",
    }
