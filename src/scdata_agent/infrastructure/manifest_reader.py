"""Read a curated manifest file into normalized rows.

Only the I/O lives here. Which columns mean what, and which roles are legal,
is dialect data in ``domain.manifest`` -- so a new spreadsheet layout is a
registration, and a new *file format* is a branch in this module alone.
"""

from __future__ import annotations

import csv
import hashlib
from pathlib import Path
from typing import Any

from scdata_agent.domain.manifest import (
    ManifestDialect,
    ManifestDocument,
    ManifestFormatError,
    ManifestIssue,
    ManifestRow,
    detect_dialect,
    dialect_by_id,
    normalize_row,
    preflight_manifest,
)

XLSX_EXTRA_HINT = (
    "XLSX manifests require the optional 'manifest' extra: pip install -e '.[manifest]'"
)


def read_tabular(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    """Return (headers, rows) for a manifest file, dispatching on suffix."""
    suffix = path.suffix.casefold()
    if suffix in {".csv", ".tsv"}:
        return _read_delimited(path, delimiter="\t" if suffix == ".tsv" else ",")
    if suffix in {".xlsx", ".xlsm"}:
        return _read_xlsx(path)
    raise ManifestFormatError(
        f"Unsupported manifest format {path.suffix!r}; expected .csv, .tsv or .xlsx."
    )


def read_manifest(path: str | Path, dialect_id: str | None = None) -> ManifestDocument:
    """Parse a manifest into rows plus every issue found while doing so."""
    source = Path(path)
    if not source.is_file():
        raise ManifestFormatError(f"Manifest not found: {source}")
    headers, raw_rows = read_tabular(source)
    dialect = _resolve_dialect(headers, dialect_id)
    rows: list[ManifestRow] = []
    issues: list[ManifestIssue] = []
    for index, raw in enumerate(raw_rows, start=1):
        row, row_issues = normalize_row(dialect, raw, index)
        issues.extend(row_issues)
        if row is not None:
            rows.append(row)
    issues.extend(preflight_manifest(rows))
    return ManifestDocument(
        source_path=str(source.resolve()),
        source_sha256=file_sha256(source),
        dialect_id=dialect.dialect_id,
        headers=headers,
        rows=rows,
        issues=issues,
    )


def file_sha256(path: Path) -> str:
    """Content hash of the manifest, so a claim can name the file it came from."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _resolve_dialect(headers: list[str], dialect_id: str | None) -> ManifestDialect:
    if dialect_id is not None:
        dialect = dialect_by_id(dialect_id)
        missing = [item for item in dialect.required_headers() if item not in set(headers)]
        if missing:
            raise ManifestFormatError(
                f"Dialect {dialect_id!r} expects column(s) not present: {', '.join(missing)}."
            )
        return dialect
    detected = detect_dialect(headers)
    if detected is None:
        raise ManifestFormatError(
            "Could not identify the manifest dialect from its headers; pass --dialect "
            "explicitly. Guessing would silently mismap every column."
        )
    return detected


def _read_delimited(path: Path, *, delimiter: str) -> tuple[list[str], list[dict[str, str]]]:
    # utf-8-sig: the source manifests are written by Excel and carry a BOM,
    # which would otherwise become part of the first column's header.
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=delimiter)
        headers = [item.strip() for item in (reader.fieldnames or [])]
        if not headers:
            raise ManifestFormatError(f"Manifest has no header row: {path}")
        rows = [{key.strip(): (value or "") for key, value in row.items() if key} for row in reader]
    return headers, rows


def _read_xlsx(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    try:
        from openpyxl import load_workbook
    except ModuleNotFoundError as error:  # pragma: no cover - exercised via the message
        raise ManifestFormatError(XLSX_EXTRA_HINT) from error
    workbook = load_workbook(filename=path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        iterator = sheet.iter_rows(values_only=True)
        header_row = next(iterator, None)
        if header_row is None:
            raise ManifestFormatError(f"Manifest has no header row: {path}")
        headers = [_text(cell) for cell in header_row]
        rows = [
            {header: _text(cell) for header, cell in zip(headers, values, strict=False) if header}
            for values in iterator
            if any(cell is not None for cell in values)
        ]
    finally:
        workbook.close()
    return [item for item in headers if item], rows


def _text(cell: Any) -> str:
    if cell is None:
        return ""
    return str(cell).strip()
